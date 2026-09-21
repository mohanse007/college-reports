"""
Excel Styler Engine for College Reports
Applies standardized fonts, alignments, borders, and column widths.
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Common Borders
THIN_BORDER_SIDE = Side(style='thin', color='000000')
CELL_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE
)

# Common Fonts
FONT_TITLE = Font(name='Calibri', size=16, bold=True)
FONT_DEPT = Font(name='Calibri', size=14, bold=True)
FONT_TABLE_HEADER = Font(name='Calibri', size=11, bold=True)
FONT_DATA = Font(name='Calibri', size=11, bold=False)
FONT_BOLD_DATA = Font(name='Calibri', size=11, bold=True)

# Common Alignments
ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_LEFT = Alignment(horizontal='left', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')


def style_merged_range(ws, start_row, start_col, end_row, end_col, font=None, alignment=None, fill=None, border=None):
    """Safely apply styles to a merged range of cells in openpyxl."""
    ws.merge_cells(
        start_row=start_row, start_column=start_col,
        end_row=end_row, end_column=end_col
    )
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            if font:
                cell.font = font
            if alignment:
                cell.alignment = alignment
            if fill:
                cell.fill = fill
            if border:
                cell.border = border


def apply_table_header(ws, row_idx, headers):
    """Format a table header row."""
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=text)
        cell.font = FONT_TABLE_HEADER
        cell.alignment = ALIGN_LEFT
        cell.border = CELL_BORDER


def apply_data_row(ws, row_idx, values, bold_col1=False):
    """Format a data row with standard borders and fonts."""
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.font = FONT_BOLD_DATA if (bold_col1 and col_idx == 1) else FONT_DATA
        cell.alignment = ALIGN_CENTER if col_idx == 4 else ALIGN_LEFT
        cell.border = CELL_BORDER


def autofit_column_widths(ws, min_widths=None):
    """Auto-adjust column widths with optional minimum widths."""
    min_widths = min_widths or {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Ignore merged title rows for column width calculation to prevent extra wide columns
            if cell.row <= 4:
                continue
            if cell.value is not None:
                val_str = str(cell.value)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        calc_width = max(max_len + 3, min_widths.get(col_letter, 12))
        ws.column_dimensions[col_letter].width = calc_width
