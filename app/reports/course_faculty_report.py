"""
Course List with Faculty Details Report Generator
Converts raw course allocation Excel exports into a standardized college report.
"""
import os
import re
import html
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from .base_report import BaseReport
from ..core.faculty_manager import FacultyManager
from ..core.excel_styler import (
    style_merged_range,
    apply_table_header,
    apply_data_row,
    autofit_column_widths,
    FONT_TITLE,
    FONT_DEPT,
    FONT_TABLE_HEADER,
    FONT_DATA,
    CELL_BORDER,
    ALIGN_LEFT,
    ALIGN_CENTER
)


def get_ordinal_suffix(n: int) -> str:
    """Returns ordinal string for a number, e.g. 1 -> 1st, 2 -> 2nd, 3 -> 3rd, etc."""
    if 11 <= (n % 100) <= 13:
        return f"{n}th"
    suffix = {1: "st", 2: "nd", 3: "rd"}.get(n % 10, "th")
    return f"{n}{suffix}"


def clean_text(text: Any) -> str:
    """Cleans HTML entities, unicode dashes, and trims whitespace."""
    if text is None:
        return ""
    s = str(text)
    # Replace en-dash, em-dash, and unicode replacement chars
    s = s.replace('\u2013', '-').replace('\u2014', '-').replace('\ufffd', '').replace('', '')
    # Normalize &Amp; / &amp;
    s = re.sub(r'&amp;', '&', s, flags=re.IGNORECASE)
    s = html.unescape(s)
    # Normalize internal spaces
    s = re.sub(r'[ \t]+', ' ', s).strip()
    return s


def format_course_name(name: str) -> str:
    """Standardizes (T) and (P) suffixes on course names."""
    cleaned = clean_text(name)
    cleaned = re.sub(r'[\s\-]+[\(\[]?\s*[Tt]\s*[\)\]]?$', ' (T)', cleaned)
    cleaned = re.sub(r'[\s\-]+[\(\[]?\s*[Pp]\s*[\)\]]?$', ' (P)', cleaned)
    return cleaned


# Standard canonical department display names and ordering
CANONICAL_DEPTS = [
    "Languages",
    "Botany",
    "Agriculture",
    "BBA",
    "Chemistry",
    "B.COM",
    "BCA",
    "CS",
    "Maths",
    "MB",
    "SPL ENG",
    "Zoology",
    "PHYSICS"
]

# Exact mapping from batch/section (year-stripped, lowercased) to department heading
BATCH_TO_DEPT_MAP = {
    "bsc botany": "Botany",
    "bsc agriculture and rural development": "Agriculture",
    "bba": "BBA",
    "bsc chemistry": "Chemistry",
    "b.com computer application": "B.COM",
    "bcom computer application": "B.COM",
    "bca": "BCA",
    "bsc computer science": "CS",
    "bsc mathematics": "Maths",
    "bsc microbiology": "MB",
    "ba special english": "SPL ENG",
    "bsc zoology": "Zoology",
    "bsc physics": "PHYSICS",
}

LANGUAGE_PREFIXES = ("ENG", "TEL", "HIN", "SAN", "SKT", "FRE", "URD")


def is_language_course(code: str) -> bool:
    """Checks if a course is a general language paper (English, Telugu, Hindi, etc.)."""
    if not code:
        return False
    code_u = code.strip().upper()
    return any(code_u.startswith(p) for p in LANGUAGE_PREFIXES)


# Subject prefixes for Sheet 2 mapping
PREFIX_TO_SUBJECT = {
    "BSB": "Botany",
    "BBA": "BBA",
    "BSC": "Chemistry",
    "BC": "B COM",
    "BCM": "B COM",
    "BCA": "BCA",
    "BSCS": "Computer Science",
    "BSM": "Maths",
    "BSMB": "Microbiology",
    "BA": "Spl Eng",
    "BSZ": "Zoology",
    "BSP": "Physics",
    "BSBT": "Bio-Technology",
    "BAH": "History"
}


class CourseFacultyReport(BaseReport):
    """Generates the Department-wise Course List with Faculty Details Report."""

    @property
    def report_id(self) -> str:
        return "course_faculty_report"

    @property
    def display_name(self) -> str:
        return "Course List with Faculty Details"

    @property
    def description(self) -> str:
        return "Generates department-wise course allocation report with theory/practical pairing and staff contact numbers."

    def extract_metadata(self, raw_file_path: str) -> Dict[str, Any]:
        """Auto-detects college name, semester title, and admitted batch year."""
        college_name = "St.Ann's College for Women (A)"
        semester_title = "4th Semester Course Code, Course name with faculty Details"
        admitted_batch = "Admitted Batch 2024-2025"

        try:
            wb = openpyxl.load_workbook(raw_file_path, data_only=True)
            ws = wb.active

            # Search top 10 rows for metadata text
            meta_str = ""
            for r in range(1, min(10, ws.max_row + 1)):
                val = str(ws.cell(row=r, column=1).value or "")
                if "Term:" in val or "Batch Start Year:" in val:
                    meta_str = val
                    break

            if meta_str:
                # 1. Parse Term e.g. "Term: S5" or "Term: S4"
                term_match = re.search(r'Term:\s*S?(\d+)', meta_str, re.IGNORECASE)
                if term_match:
                    sem_num = int(term_match.group(1))
                    ord_sem = get_ordinal_suffix(sem_num)
                    semester_title = f"{ord_sem} Semester Course Code, Course name with faculty Details"

                # 2. Parse Batch Start Year e.g. "Batch Start Year: 2024"
                year_match = re.search(r'Batch Start Year:\s*(\d{4})', meta_str)
                if year_match:
                    start_year = int(year_match.group(1))
                    admitted_batch = f"Admitted Batch {start_year}-{start_year + 1}"

        except Exception as e:
            print(f"Warning extracting metadata: {e}")

        return {
            "college_name": college_name,
            "semester_title": semester_title,
            "admitted_batch": admitted_batch
        }

    def _map_batch_to_dept(self, batch_name: str) -> str:
        """Maps batch/section string to standardized department name."""
        if not batch_name:
            return ""
        # Strip trailing 4-digit years (e.g. 2024, 2025)
        cleaned = re.sub(r'\s*\b\d{4}\b.*$', '', str(batch_name)).strip().lower()
        if cleaned in BATCH_TO_DEPT_MAP:
            return BATCH_TO_DEPT_MAP[cleaned]

        # Fallback for future batches not yet in dictionary
        fb = re.sub(r'^(bsc|ba|bcom|bba|bca)\s+', '', cleaned).strip()
        return fb.title() if fb else str(batch_name).strip()


    def _parse_raw_data(self, raw_file_path: str) -> List[Dict[str, Any]]:
        """Reads raw data, forward-fills multi-batch rows, and filters out non-academic rows."""
        wb = openpyxl.load_workbook(raw_file_path, data_only=True)
        ws = wb.active

        # Find header row
        header_row = 5
        for r in range(1, 15):
            val = str(ws.cell(row=r, column=2).value or "").strip().lower()
            if "course code" in val:
                header_row = r
                break

        records = []
        last_course_code = None
        last_course_name = None
        last_assigned_faculty = None
        last_course_community = None

        for r in range(header_row + 1, ws.max_row + 1):
            c_code = clean_text(ws.cell(row=r, column=2).value)
            c_name = clean_text(ws.cell(row=r, column=3).value)
            c_comm = clean_text(ws.cell(row=r, column=4).value)
            faculty = clean_text(ws.cell(row=r, column=5).value)
            batch = clean_text(ws.cell(row=r, column=6).value)

            if not batch:
                continue

            # Update or forward-fill course info
            if c_code:
                last_course_code = c_code
                last_course_name = c_name
                last_course_community = c_comm
                last_assigned_faculty = faculty
            else:
                # Multi-batch row inherits course info
                c_code = last_course_code
                if not c_name:
                    c_name = last_course_name
                if not faculty:
                    faculty = last_assigned_faculty
                if not c_comm:
                    c_comm = last_course_community

            # Handle courses where code is blank but community has it (e.g. CVAC S551, SD S305, etc.)
            if not c_code and c_comm:
                comm_parts = c_comm.split('-', 1)
                if len(comm_parts) == 2:
                    possible_code = comm_parts[0].strip().upper()
                    if any(possible_code.startswith(p) for p in ["CVAC", "SDS", "SD", "MD", "ENG", "TEL", "HIN", "SAN", "SKT"]):
                        c_code = possible_code
                        c_name = comm_parts[1].strip()

            # Filter out activities / non-academic rows
            is_activity = (
                c_code.strip().upper() in ["GAMES", "LIBRARY", "MI"] or
                c_name.strip().upper() in ["GAMES", "LIBRARY", "MI"] or
                c_comm.strip().upper() in ["GAMES - GAMES", "LIBRARY - LIBRARY", "MI - MI"]
            )
            if is_activity:
                continue

            if not c_code or not c_name:
                continue

            # Route language courses to Languages department
            if is_language_course(c_code):
                dept_name = "Languages"
            else:
                dept_name = self._map_batch_to_dept(batch)

            records.append({
                "course_code": c_code,
                "course_name": format_course_name(c_name),
                "faculty": faculty,
                "batch": batch,
                "dept": dept_name
            })

        return records

    def _pair_courses_for_dept(self, dept_records: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """
        Organizes courses so that practical courses (starting with 'P ') are paired
        immediately after their corresponding theory courses.
        Major courses appear first, followed by Minor courses, followed by Skill/CVAC courses.
        """
        theory_list = []
        practical_dict = {}

        # De-duplicate by (course_code, faculty)
        seen = set()
        for rec in dept_records:
            key = (rec["course_code"].upper(), rec["faculty"].upper())
            if key in seen:
                continue
            seen.add(key)

            code = rec["course_code"]
            if code.upper().startswith("P "):
                base = code[2:].strip().upper()
                practical_dict[base] = rec
            else:
                theory_list.append(rec)

        def get_course_sort_key(rec):
            code_u = rec["course_code"].strip().upper()
            # Category rank:
            # Rank 1: Major courses
            # Rank 2: Minor courses (ending with 'M')
            # Rank 3: Skill / Multi-Disciplinary / Value Added courses (SD, SDS, MD, CVAC, SEC)
            if any(code_u.startswith(p) for p in ["SDS", "SD ", "SD", "MD", "CVAC", "SEC"]):
                cat = 3
            elif code_u.endswith("M"):
                cat = 2
            else:
                cat = 1

            # Extract numeric digits (e.g. S351 -> 351, 305 -> 305)
            num_match = re.search(r'(\d+)', code_u)
            num = int(num_match.group(1)) if num_match else 999

            # Extract alphabetic prefix (e.g. BSC, BSCS, BSM, BSP, SD, MD)
            prefix_match = re.match(r'^([A-Z]+)', code_u)
            prefix = prefix_match.group(1) if prefix_match else ""

            return (cat, num, prefix, code_u)

        # Sort theory courses: Category -> Numeric code order -> Prefix
        theory_list.sort(key=get_course_sort_key)

        ordered = []
        for t_rec in theory_list:
            ordered.append(t_rec)
            t_base = t_rec["course_code"].strip().upper()
            if t_base in practical_dict:
                p_rec = practical_dict.pop(t_base)
                ordered.append(p_rec)

        # Any remaining practical courses without a theory match
        for remaining_p in practical_dict.values():
            ordered.append(remaining_p)

        return ordered

    def generate(self, raw_file_path: str, output_dir: str, options: Dict[str, Any]) -> str:
        """Generates the formatted Excel report."""
        os.makedirs(output_dir, exist_ok=True)

        # Load metadata
        meta = self.extract_metadata(raw_file_path)
        college_name = options.get("college_name") or meta["college_name"]
        semester_title = options.get("semester_title") or meta["semester_title"]
        admitted_batch = options.get("admitted_batch") or meta["admitted_batch"]

        # Faculty lookup
        faculty_mgr = FacultyManager()

        # Parse records
        records = self._parse_raw_data(raw_file_path)

        # Group records by department
        dept_records_map = {}
        for r in records:
            dept = r["dept"]
            if dept not in dept_records_map:
                dept_records_map[dept] = []
            dept_records_map[dept].append(r)

        # Create workbook
        wb = openpyxl.Workbook()
        
        # ==========================================
        # Sheet 1: Department-wise Course List
        # ==========================================
        ws1 = wb.active
        ws1.title = "Sheet1"

        # College Headers (Rows 1-3)
        ws1.cell(row=1, column=1, value=college_name)
        style_merged_range(ws1, 1, 1, 1, 4, font=FONT_TITLE, alignment=ALIGN_CENTER)

        ws1.cell(row=2, column=1, value=semester_title)
        style_merged_range(ws1, 2, 1, 2, 4, font=FONT_TITLE, alignment=ALIGN_CENTER)

        ws1.cell(row=3, column=1, value=admitted_batch)
        style_merged_range(ws1, 3, 1, 3, 4, font=FONT_TITLE, alignment=ALIGN_CENTER)

        current_row = 4

        # Order departments according to canonical order
        ordered_depts = []
        for d_name in CANONICAL_DEPTS:
            if d_name in dept_records_map:
                ordered_depts.append(d_name)
        for d_name in dept_records_map:
            if d_name not in ordered_depts:
                ordered_depts.append(d_name)

        table_headers = ["Course Code", "Course Name", "Name of Staff", "Phone Number"]

        for dept_name in ordered_depts:
            dept_recs = dept_records_map[dept_name]
            paired_recs = self._pair_courses_for_dept(dept_recs)

            # Department Title Row
            ws1.cell(row=current_row, column=1, value=dept_name)
            style_merged_range(ws1, current_row, 1, current_row, 2, font=FONT_DEPT, alignment=ALIGN_LEFT)
            current_row += 1

            # Table Header Row
            apply_table_header(ws1, current_row, table_headers)
            current_row += 1

            # Data Rows
            for item in paired_recs:
                staff_name = item["faculty"]
                phone = faculty_mgr.get_phone(staff_name)
                c_name = item["course_name"]

                apply_data_row(
                    ws1,
                    current_row,
                    [item["course_code"], c_name, staff_name, phone],
                    bold_col1=False
                )
                current_row += 1

            # Empty separator row
            current_row += 1

        autofit_column_widths(ws1, min_widths={"A": 16, "B": 55, "C": 30, "D": 18})

        # ==========================================
        # Sheet 2: Master Course List
        # ==========================================
        ws2 = wb.create_sheet(title="Sheet2")
        s2_headers = ["Sl.No", "Subject", "Subject Code", "Title of the paper"]
        apply_table_header(ws2, 1, s2_headers)

        # Categorize unique courses
        unique_courses = {}
        for r in records:
            c_code = r["course_code"]
            if c_code.upper().startswith("P "):
                continue # Sheet2 only lists main/theory papers
            if c_code not in unique_courses:
                unique_courses[c_code] = {
                    "code": c_code,
                    "name": r["course_name"],
                    "dept": r["dept"]
                }

        lang_courses = []
        skill_courses = []
        multi_courses = []
        major_courses = []
        minor_courses = []

        for c_info in unique_courses.values():
            code_u = c_info["code"].upper()
            if is_language_course(code_u):
                lang_courses.append(c_info)
            elif code_u.startswith("SDS") or code_u.startswith("SD"):
                skill_courses.append(c_info)
            elif code_u.startswith("MD") or code_u.startswith("CVAC"):
                multi_courses.append(c_info)
            elif c_info["code"].endswith("m") or c_info["code"].endswith("M"):
                minor_courses.append(c_info)
            else:
                major_courses.append(c_info)

        s2_row = 2
        sl_no = 1

        def append_section(title, courses, is_subject_dept=False):
            nonlocal s2_row, sl_no
            if not courses:
                return
            if title:
                # Section heading
                ws2.cell(row=s2_row, column=1, value=title).font = FONT_TABLE_HEADER
                for c in range(1, 5):
                    ws2.cell(row=s2_row, column=c).border = CELL_BORDER
                s2_row += 1

            curr_subj = None
            for c in courses:
                subj = title if not is_subject_dept else c["dept"]
                subj_disp = subj if subj != curr_subj else ""
                curr_subj = subj

                apply_data_row(
                    ws2,
                    s2_row,
                    [sl_no, subj_disp, c["code"], c["name"]]
                )
                ws2.cell(row=s2_row, column=1).alignment = ALIGN_CENTER
                sl_no += 1
                s2_row += 1

        append_section("Languages", lang_courses)
        append_section("Skill Course", skill_courses)
        append_section("Multi Disciplinary", multi_courses)
        append_section("Major", major_courses, is_subject_dept=True)
        append_section("Minor", minor_courses, is_subject_dept=True)

        autofit_column_widths(ws2, min_widths={"A": 10, "B": 25, "C": 18, "D": 55})

        # Save output file
        raw_basename = os.path.splitext(os.path.basename(raw_file_path))[0]
        # Clean up output filename nicely based on semester title
        sem_clean = re.sub(r'[^a-zA-Z0-9]', '_', semester_title.split(',')[0]).strip('_')
        out_filename = f"{sem_clean}_Course_List_with_Faculty_Details.xlsx"
        output_file_path = os.path.join(output_dir, out_filename)
        try:
            wb.save(output_file_path)
        except PermissionError:
            import datetime
            ts = datetime.datetime.now().strftime("%H%M%S")
            base, ext = os.path.splitext(output_file_path)
            output_file_path = f"{base}_{ts}{ext}"
            wb.save(output_file_path)

        return output_file_path
